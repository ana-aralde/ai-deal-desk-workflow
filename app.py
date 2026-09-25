from __future__ import annotations

import json
import os
import re
import uuid
import hmac
from io import BytesIO
from dataclasses import dataclass, asdict
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Tuple

import pandas as pd
import streamlit as st

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None

try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None

try:
    from openai import OpenAI
except Exception:  # pragma: no cover
    OpenAI = None


APP_TITLE = "AI-Assisted Deal Desk Intake, Routing & QA Workflow"
APP_VERSION = "5.5"


# -----------------------------
# Configuration helpers
# -----------------------------
def secret(name: str, default: Any = None) -> Any:
    """Read from Streamlit secrets first, then environment variables."""
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:
        pass
    return os.getenv(name, default)


def bool_secret(name: str, default: bool = False) -> bool:
    value = str(secret(name, str(default))).strip().lower()
    return value in {"1", "true", "yes", "y", "on"}


def int_secret(name: str, default: int) -> int:
    try:
        return int(secret(name, default))
    except Exception:
        return default


# -----------------------------
# Privacy / data minimization
# -----------------------------
PII_PATTERNS = [
    (re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I), "[EMAIL_REDACTED]"),
    (re.compile(r"(?<!\d)(?:\+?55\s*)?(?:\(?\d{2}\)?\s*)?9?\d{4}[-\s]?\d{4}(?!\d)"), "[PHONE_REDACTED]"),
    (re.compile(r"(?<!\d)\d{3}\.?\d{3}\.?\d{3}-?\d{2}(?!\d)"), "[CPF_REDACTED]"),
    (re.compile(r"(?<!\d)\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}(?!\d)"), "[CNPJ_REDACTED]"),
    (re.compile(r"(?<!\d)(?:\d[ -]*?){13,19}(?!\d)"), "[LONG_NUMBER_REDACTED]"),
]


def redact_common_pii(text: str) -> Tuple[str, int]:
    redacted = text or ""
    count = 0
    for pattern, replacement in PII_PATTERNS:
        redacted, n = pattern.subn(replacement, redacted)
        count += n
    return redacted, count


def truncate(text: str, max_chars: int = 6000) -> str:
    text = text or ""
    return text if len(text) <= max_chars else text[:max_chars] + "\n[TRUNCATED]"


# -----------------------------
# Domain rules
# -----------------------------
SME_MAP = {
    "Tender / RFP / Procurement Questionnaire": ["Deal Desk"],
    "Security Questionnaire": ["Security / GRC"],
    "Privacy / DPA": ["Privacy", "Legal"],
    "Contract / Legal Terms": ["Legal"],
    "Commercial / Pricing": ["Finance", "Sales"],
    "Vendor Onboarding": ["Deal Desk", "Finance"],
    "Attestation / Representation": ["Legal", "Deal Desk"],
    "Product Capability": ["Product"],
    "Infrastructure / Data Residency": ["Infrastructure", "Privacy"],
    "Other": ["Deal Desk"],
}

HIGH_RISK_CUES = [
    "data residency", "cross-border", "dpa", "privacy", "gdpr", "dora",
    "fedramp", "soc 2", "iso 27001", "iso 27701", "cyber essentials",
    "indemnity", "liability", "warranty", "audit right", "subprocessor",
    "certification", "roadmap", "guarantee", "breach", "security incident",
    "sanctions", "aml", "cft", "regulator", "government", "federal",
]

MEDIUM_RISK_CUES = [
    "insurance", "payment", "pricing", "renewal", "sla", "service level",
    "entity", "registration", "attestation", "representation", "tax",
]


def deterministic_triage(request_type: str, text: str, deadline: date) -> Dict[str, Any]:
    normalized = (text or "").lower()
    days_left = (deadline - date.today()).days
    high_hits = [cue for cue in HIGH_RISK_CUES if cue in normalized]
    med_hits = [cue for cue in MEDIUM_RISK_CUES if cue in normalized]

    if request_type in {"Privacy / DPA", "Contract / Legal Terms", "Security Questionnaire", "Infrastructure / Data Residency"}:
        review_path = "Enhanced"
    elif high_hits:
        review_path = "Enhanced"
    elif med_hits or request_type in {"Commercial / Pricing", "Attestation / Representation"}:
        review_path = "Standard"
    else:
        review_path = "Simplified"

    if days_left < 0:
        urgency = "Overdue"
    elif days_left <= 1:
        urgency = "Critical"
    elif days_left <= 3:
        urgency = "High"
    elif days_left <= 7:
        urgency = "Medium"
    else:
        urgency = "Normal"

    owners = list(SME_MAP.get(request_type, ["Deal Desk"]))
    if any(x in normalized for x in ["privacy", "dpa", "gdpr", "data residency", "cross-border"]):
        owners = sorted(set(owners + ["Privacy", "Legal"]))
    if any(x in normalized for x in ["security", "soc 2", "iso 27001", "fedramp", "cyber"]):
        owners = sorted(set(owners + ["Security / GRC"]))
    if any(x in normalized for x in ["pricing", "payment", "insurance"]):
        owners = sorted(set(owners + ["Finance"]))
    if any(x in normalized for x in ["roadmap", "capability", "feature"]):
        owners = sorted(set(owners + ["Product"]))

    return {
        "review_path": review_path,
        "urgency": urgency,
        "days_left": days_left,
        "suggested_owners": owners,
        "risk_cues": high_hits + med_hits,
    }


def keyword_relevance(query: str, source_text: str) -> float:
    stop = {"the", "and", "or", "of", "to", "a", "in", "for", "is", "are", "on", "with", "de", "e", "o", "a", "do", "da", "para"}
    q = {w for w in re.findall(r"[A-Za-zÀ-ÿ0-9_-]+", query.lower()) if len(w) > 2 and w not in stop}
    s = {w for w in re.findall(r"[A-Za-zÀ-ÿ0-9_-]+", source_text.lower()) if len(w) > 2 and w not in stop}
    if not q:
        return 0.0
    return len(q & s) / max(1, len(q))


def retrieve_sources(query: str, sources: List[Dict[str, Any]], limit: int = 4) -> List[Dict[str, Any]]:
    approved = [s for s in sources if s.get("status") == "Approved"]
    scored = [
        (keyword_relevance(query, " ".join([s.get("title", ""), s.get("content", ""), s.get("owner", "")])), s)
        for s in approved
    ]
    scored.sort(key=lambda x: x[0], reverse=True)
    # Conservative lexical retrieval for the public demo: do not fall back to unrelated Approved sources.
    # A production system would normally use governed semantic retrieval + metadata filters.
    return [s for score, s in scored[:limit] if score >= 0.05]




# -----------------------------
# Request-document ingestion
# -----------------------------
ALLOWED_UPLOAD_TYPES = ["pdf", "docx", "xlsx", "csv", "txt"]
MAX_UPLOAD_FILES = 5
MAX_UPLOAD_FILE_BYTES = 5 * 1024 * 1024  # 5 MB per file in the public demo
MAX_REQUIREMENTS_PER_REQUEST = 200
MAX_ATTACHMENT_TEXT_CHARS = 40000

REQUIREMENT_COLUMN_HINTS = [
    "requirement", "requirements", "question", "questions", "control", "controls",
    "description", "criteria", "criterion", "item", "request", "requirement text",
]


def normalize_space(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def requirement_category(text: str) -> Tuple[str, str]:
    t = (text or "").lower()
    if any(k in t for k in ["data residency", "hosting region", "hosting location", "hosted", "infrastructure", "cloud region"]):
        return "Infrastructure / Data Residency", "Infrastructure + Privacy"
    if any(k in t for k in ["privacy", "gdpr", "dpa", "personal data", "cross-border", "data transfer", "subprocessor"]):
        return "Privacy", "Privacy / Legal"
    if any(k in t for k in ["security", "soc 2", "iso 27001", "fedramp", "cyber", "encryption", "vulnerability", "incident response", "access control"]):
        return "Security / GRC", "Security / GRC"
    if any(k in t for k in ["indemnity", "liability", "warranty", "governing law", "contract", "legal terms", "audit right", "termination"]):
        return "Legal", "Legal"
    if any(k in t for k in ["pricing", "price", "payment", "insurance", "tax", "renewal", "discount", "invoice"]):
        return "Commercial / Finance", "Finance"
    if any(k in t for k in ["roadmap", "capability", "feature", "integration", "functionality", "product"]):
        return "Product", "Product"
    return "General / Deal Desk", "Deal Desk"


def make_requirement(text: str, source_file: str, source_location: str = "") -> Dict[str, Any]:
    cleaned = normalize_space(re.sub(r"^[\s\-–—•▪◦*]+", "", text or ""))
    category, owner = requirement_category(cleaned)
    return {
        "requirement_id": f"R-{uuid.uuid4().hex[:7].upper()}",
        "requirement": cleaned[:1200],
        "category": category,
        "suggested_owner": owner,
        "source_file": source_file,
        "source_location": source_location,
    }


def requirement_candidates_from_text(text: str, source_file: str, location_prefix: str = "") -> List[Dict[str, Any]]:
    """Extract conservative line/paragraph candidates without using AI."""
    out: List[Dict[str, Any]] = []
    seen = set()
    for idx, raw in enumerate((text or "").splitlines(), start=1):
        line = normalize_space(raw)
        if not line or len(line) < 12:
            continue
        # Skip obvious page-only / decorative fragments.
        if re.fullmatch(r"(?:page\s*)?\d+(?:\s*/\s*\d+)?", line, re.I):
            continue
        # Remove common enumeration prefixes while preserving the substance.
        cleaned = normalize_space(re.sub(r"^(?:\(?\d+[.)]|[A-Za-z][.)]|[-–—•▪◦*])\s*", "", line))
        if len(cleaned) < 12:
            continue
        key = cleaned.lower()
        if key in seen:
            continue
        seen.add(key)
        location = f"{location_prefix}line {idx}" if location_prefix else f"line {idx}"
        out.append(make_requirement(cleaned, source_file, location))
        if len(out) >= MAX_REQUIREMENTS_PER_REQUEST:
            break
    return out


def dataframe_requirement_rows(df: pd.DataFrame, source_file: str, sheet_name: str = "") -> List[Dict[str, Any]]:
    if df is None or df.empty:
        return []
    df = df.copy().dropna(how="all")
    df.columns = [normalize_space(c) or f"Column {i+1}" for i, c in enumerate(df.columns)]
    hint_col = None
    for col in df.columns:
        c = col.lower()
        if any(h == c or h in c for h in REQUIREMENT_COLUMN_HINTS):
            hint_col = col
            break

    out: List[Dict[str, Any]] = []
    for ridx, row in df.iterrows():
        if hint_col is not None:
            text = normalize_space(row.get(hint_col, ""))
        else:
            parts = []
            for col, val in row.items():
                if pd.isna(val):
                    continue
                v = normalize_space(val)
                if v:
                    parts.append(f"{col}: {v}")
            text = " | ".join(parts)
        if len(text) < 8:
            continue
        loc = f"{sheet_name + ' · ' if sheet_name else ''}row {int(ridx) + 2 if isinstance(ridx, int) else ridx}"
        out.append(make_requirement(text, source_file, loc))
        if len(out) >= MAX_REQUIREMENTS_PER_REQUEST:
            break
    return out


def extract_uploaded_document(uploaded_file: Any) -> Dict[str, Any]:
    """Parse a supported upload in memory. No raw file is persisted by the app."""
    name = getattr(uploaded_file, "name", "uploaded-file")
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    raw = uploaded_file.getvalue()
    result = {
        "name": name,
        "extension": ext,
        "size_bytes": len(raw),
        "text": "",
        "requirements": [],
        "warnings": [],
    }
    if len(raw) > MAX_UPLOAD_FILE_BYTES:
        result["warnings"].append(f"Skipped: file exceeds the {MAX_UPLOAD_FILE_BYTES // (1024*1024)} MB public-demo limit.")
        return result

    try:
        if ext == "pdf":
            if PdfReader is None:
                raise RuntimeError("PDF parser is unavailable. Check requirements.txt / deployment dependencies.")
            reader = PdfReader(BytesIO(raw))
            pages = []
            reqs = []
            for pnum, page in enumerate(reader.pages[:80], start=1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    pages.append(f"[Page {pnum}]\n{page_text}")
                    reqs.extend(requirement_candidates_from_text(page_text, name, f"page {pnum} · "))
            result["text"] = "\n\n".join(pages)
            result["requirements"] = reqs[:MAX_REQUIREMENTS_PER_REQUEST]
            if not result["text"].strip():
                result["warnings"].append("No machine-readable PDF text was found. OCR is intentionally not enabled in this public demo.")

        elif ext == "docx":
            if Document is None:
                raise RuntimeError("DOCX parser is unavailable. Check requirements.txt / deployment dependencies.")
            doc = Document(BytesIO(raw))
            blocks = []
            reqs = []
            for idx, para in enumerate(doc.paragraphs, start=1):
                t = para.text.strip()
                if t:
                    blocks.append(t)
                    reqs.extend(requirement_candidates_from_text(t, name, f"paragraph {idx} · "))
            for tidx, table in enumerate(doc.tables, start=1):
                for ridx, row in enumerate(table.rows, start=1):
                    cells = [normalize_space(c.text) for c in row.cells]
                    row_text = " | ".join([c for c in cells if c])
                    if row_text:
                        blocks.append(row_text)
                        reqs.append(make_requirement(row_text, name, f"table {tidx} · row {ridx}"))
            result["text"] = "\n".join(blocks)
            result["requirements"] = reqs[:MAX_REQUIREMENTS_PER_REQUEST]

        elif ext == "xlsx":
            book = pd.ExcelFile(BytesIO(raw), engine="openpyxl")
            text_parts = []
            reqs = []
            for sheet in book.sheet_names[:20]:
                df = pd.read_excel(book, sheet_name=sheet, dtype=str).fillna("")
                if df.empty:
                    continue
                text_parts.append(f"[Sheet: {sheet}]\n" + df.astype(str).head(250).to_csv(index=False))
                reqs.extend(dataframe_requirement_rows(df.head(250), name, sheet))
                if len(reqs) >= MAX_REQUIREMENTS_PER_REQUEST:
                    break
            result["text"] = "\n\n".join(text_parts)
            result["requirements"] = reqs[:MAX_REQUIREMENTS_PER_REQUEST]

        elif ext == "csv":
            last_error = None
            df = None
            for encoding in ("utf-8-sig", "utf-8", "latin-1"):
                try:
                    df = pd.read_csv(BytesIO(raw), dtype=str, encoding=encoding).fillna("")
                    break
                except Exception as exc:
                    last_error = exc
            if df is None:
                raise last_error or RuntimeError("Unable to parse CSV.")
            result["text"] = df.head(500).to_csv(index=False)
            result["requirements"] = dataframe_requirement_rows(df.head(500), name)[:MAX_REQUIREMENTS_PER_REQUEST]

        elif ext == "txt":
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("latin-1", errors="replace")
            result["text"] = text
            result["requirements"] = requirement_candidates_from_text(text, name)[:MAX_REQUIREMENTS_PER_REQUEST]

        else:
            result["warnings"].append("Unsupported file type.")
    except Exception as exc:
        result["warnings"].append(f"Could not parse this file: {exc}")

    result["text"] = truncate(result["text"], MAX_ATTACHMENT_TEXT_CHARS)
    return result


def combined_request_text(request: Dict[str, Any], max_chars: int = 14000) -> str:
    description = request.get("description", "") or ""
    attachment_text = request.get("attachment_text", "") or ""
    req_lines = [r.get("requirement", "") for r in request.get("requirements", [])[:80] if r.get("requirement")]
    combined = description
    if attachment_text:
        combined += "\n\n[Uploaded request files]\n" + attachment_text
    if req_lines:
        combined += "\n\n[Extracted requirements]\n" + "\n".join(f"- {x}" for x in req_lines)
    return truncate(combined, max_chars)


def best_candidate_source(requirement_text: str, sources: List[Dict[str, Any]]) -> Tuple[str, float]:
    best_id, best_score = "—", 0.0
    for s in sources:
        if s.get("status") != "Approved":
            continue
        haystack = " ".join([s.get("title", ""), s.get("content", ""), s.get("owner", "")])
        score = keyword_relevance(requirement_text, haystack)
        if score > best_score:
            best_id, best_score = s.get("source_id") or "—", score
    # Conservative threshold: this is a candidate match, never proof that the source answers the requirement.
    return (best_id, best_score) if best_score >= 0.12 else ("—", best_score)


# -----------------------------
# AI service
# -----------------------------
def ai_available() -> bool:
    return bool(secret("OPENAI_API_KEY")) and OpenAI is not None


def ai_permitted() -> bool:
    """Allow paid AI only when explicitly public or unlocked for this browser session."""
    if bool_secret("ALLOW_PUBLIC_AI", False):
        return True
    required_code = secret("AI_DEMO_CODE")
    return bool(required_code) and bool(st.session_state.get("ai_unlocked", False))


def call_ai_analysis(request: Dict[str, Any], sources: List[Dict[str, Any]]) -> Dict[str, Any]:
    if not ai_available():
        raise RuntimeError("OpenAI API key is not configured.")
    if not ai_permitted():
        raise RuntimeError("Live AI is locked for this session. Enter the demo access code in AI Review.")

    max_calls = int_secret("MAX_AI_CALLS_PER_SESSION", 3)
    used = st.session_state.get("ai_calls", 0)
    if used >= max_calls:
        raise RuntimeError(f"Session AI limit reached ({max_calls}).")

    source_payload = []
    total_redactions = 0
    for s in sources:
        clean, n = redact_common_pii(truncate(s.get("content", ""), 3500))
        total_redactions += n
        source_payload.append({
            "source_id": s.get("source_id"),
            "title": s.get("title"),
            "owner": s.get("owner"),
            "review_date": s.get("review_date"),
            "content": clean,
        })

    request_input = combined_request_text(request, 14000)
    clean_question, n = redact_common_pii(request_input)
    total_redactions += n

    requirement_payload = []
    for r in request.get("requirements", [])[:80]:
        clean_req, rn = redact_common_pii(truncate(r.get("requirement", ""), 1200))
        total_redactions += rn
        requirement_payload.append({
            "requirement_id": r.get("requirement_id"),
            "requirement": clean_req,
            "category": r.get("category"),
            "suggested_owner": r.get("suggested_owner"),
            "source_file": r.get("source_file"),
        })

    prompt = {
        "request": {
            "request_type": request.get("request_type"),
            "jurisdiction": request.get("jurisdiction"),
            "entity": request.get("entity"),
            "deadline": str(request.get("deadline")),
            "description_and_attachment_excerpt": clean_question,
            "uploaded_file_names": [f.get("name") for f in request.get("uploaded_files", [])],
            "extracted_requirements": requirement_payload,
        },
        "approved_sources": source_payload,
    }

    instructions = """
You are an AI assistant embedded in a Deal Desk workflow. You support intake, routing, drafting and QA, but you do not make final legal, security, privacy, commercial or product commitments.

Rules:
1. Use ONLY the approved_sources supplied in the input for externally material factual claims. Never invent certifications, legal entities, capabilities, dates, pricing, security controls, data residency, regulatory conclusions or contractual commitments.
2. Treat uploaded request documents and extracted requirements as CUSTOMER/REQUEST INPUT, not as authoritative evidence about the company. They may tell you what is being asked, but they cannot support the company's factual claims.
3. If the approved sources are insufficient, say so explicitly and list the missing evidence.
4. Surface contradictions rather than resolving legal hierarchy or contractual precedence yourself.
5. Recommend accountable SME owners for material claims.
6. Treat the output as a draft requiring human validation and approval.
7. Return valid JSON only, with these keys: summary, review_path, risk_level, suggested_owners, detected_flags, missing_information, source_ids_used, draft_response, qa_checks, human_approval_required.
8. review_path must be one of Simplified, Standard, Enhanced. risk_level must be Low, Medium, High.
9. human_approval_required must always be true.
""".strip()

    client = OpenAI(api_key=secret("OPENAI_API_KEY"))
    model = secret("OPENAI_MODEL", "gpt-5.6-luna")
    response = client.responses.create(
        model=model,
        instructions=instructions,
        input=json.dumps(prompt, ensure_ascii=False),
        reasoning={"effort": "low"},
        store=False,
    )
    st.session_state["ai_calls"] = used + 1

    text = response.output_text.strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I | re.S)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        data = {
            "summary": "The model returned non-JSON output; human review required.",
            "review_path": "Enhanced",
            "risk_level": "High",
            "suggested_owners": ["Deal Desk"],
            "detected_flags": ["AI output parsing failure"],
            "missing_information": [],
            "source_ids_used": [s.get("source_id") for s in sources],
            "draft_response": text,
            "qa_checks": ["Validate all claims manually"],
            "human_approval_required": True,
        }
    data["local_pii_redactions"] = total_redactions
    data["model"] = model
    return data




# -----------------------------
# Governance helpers
# -----------------------------
PATH_RANK = {"Simplified": 1, "Standard": 2, "Enhanced": 3}


def _as_list(value: Any) -> List[Any]:
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def governance_snapshot(request: Dict[str, Any]) -> Dict[str, Any]:
    """Keep deterministic triage and AI recommendation separate and surface unresolved blockers."""
    baseline = request.get("review_path") or "Not set"
    result = request.get("ai_result") or {}
    ai_path = result.get("review_path")
    risk_level = result.get("risk_level")
    missing = _as_list(result.get("missing_information"))
    flags = _as_list(result.get("detected_flags"))

    high_flags = []
    for flag in flags:
        if isinstance(flag, dict):
            severity = str(flag.get("severity", "")).strip().lower()
            if severity == "high":
                high_flags.append(flag)
        elif risk_level == "High":
            # Unstructured flags inherit the overall High risk level for release-gate purposes.
            high_flags.append(flag)

    path_discrepancy = bool(ai_path and ai_path != baseline)
    blockers: List[str] = []
    if path_discrepancy:
        blockers.append(f"Triage discrepancy: deterministic baseline is {baseline}, while AI recommends {ai_path}.")
    if risk_level == "High":
        blockers.append("AI classified the request as High risk and human resolution is required before release.")
    if missing:
        blockers.append(f"AI identified {len(missing)} missing-information item(s) that require human resolution.")
    if high_flags:
        blockers.append(f"AI identified {len(high_flags)} High-severity flag(s) that require human resolution.")

    effective_path = baseline
    if ai_path and PATH_RANK.get(ai_path, 0) > PATH_RANK.get(baseline, 0):
        effective_path = ai_path

    request_status = request.get("status")
    override_approved = bool(request.get("human_override", {}).get("approved"))

    if request_status == "Released":
        if blockers and override_approved:
            governance_status = "Released with documented human decision"
        elif blockers:
            governance_status = "Released - historical blockers retained"
        elif result:
            governance_status = "Released after AI review"
        else:
            governance_status = "Released"
    elif not result:
        governance_status = "Baseline only"
    elif blockers:
        governance_status = "Escalation required"
    else:
        governance_status = "AI reviewed - no unresolved blocker"

    return {
        "baseline_path": baseline,
        "ai_path": ai_path,
        "effective_path": effective_path,
        "risk_level": risk_level,
        "missing": missing,
        "flags": flags,
        "high_flags": high_flags,
        "path_discrepancy": path_discrepancy,
        "blockers": blockers,
        "governance_status": governance_status,
    }


# -----------------------------
# State and sample data
# -----------------------------
def init_state() -> None:
    st.session_state.setdefault("requests", [])
    st.session_state.setdefault("sources", [])
    st.session_state.setdefault("ai_calls", 0)
    st.session_state.setdefault("ai_unlocked", False)


def load_demo_data() -> None:
    if st.session_state["requests"] or st.session_state["sources"]:
        return
    st.session_state["sources"] = [
        {
            "source_id": "SRC-001",
            "title": "Approved Security Statement - Fictional",
            "owner": "Security / GRC",
            "status": "Approved",
            "review_date": str(date.today() + timedelta(days=60)),
            "content": "Fictional demo source: The company maintains an information security program and requires Security/GRC approval before externally representing certification scope or control effectiveness.",
        },
        {
            "source_id": "SRC-002",
            "title": "Approved Privacy Statement - Fictional",
            "owner": "Privacy / Legal",
            "status": "Approved",
            "review_date": str(date.today() + timedelta(days=45)),
            "content": "Fictional demo source: Privacy and cross-border data transfer commitments must be validated by Privacy/Legal. The Deal Desk may reuse approved language only when the entity, service and jurisdiction match the source scope.",
        },
        {
            "source_id": "SRC-003",
            "title": "Product Capability Note - Fictional",
            "owner": "Product",
            "status": "Approved",
            "review_date": str(date.today() + timedelta(days=30)),
            "content": "Fictional demo source: Product capability and roadmap claims require Product-owner confirmation. Future features must not be represented as committed delivery dates unless specifically approved.",
        },
    ]
    st.session_state["requests"] = [
        {
            "request_id": "REQ-DEMO-01",
            "created_at": datetime.now().isoformat(timespec="minutes"),
            "customer": "Fictional European Bank",
            "entity": "EU Entity",
            "jurisdiction": "European Union",
            "request_type": "Security Questionnaire",
            "deadline": str(date.today() + timedelta(days=4)),
            "owner": "Deal Desk",
            "status": "In Review",
            "description": "Confirm current security certification scope and whether customer data can be hosted in the EU. Provide evidence and identify any claims requiring Security, Privacy or Legal validation.",
            "review_path": "Enhanced",
            "urgency": "Medium",
            "suggested_owners": ["Security / GRC", "Privacy", "Legal"],
            "risk_cues": ["certification", "data residency"],
            "uploaded_files": [],
            "attachment_text": "",
            "requirements": [
                make_requirement("Confirm current security certification scope.", "fictional-security-questionnaire.xlsx", "row 2"),
                make_requirement("Confirm whether customer data can be hosted in the EU.", "fictional-security-questionnaire.xlsx", "row 3"),
                make_requirement("Provide evidence and identify claims requiring Security, Privacy or Legal validation.", "fictional-security-questionnaire.xlsx", "row 4"),
            ],
            "ai_result": None,
            "qa": {},
            "human_override": {},
        }
    ]


# -----------------------------
# UI components
# -----------------------------
def ai_access_panel() -> bool:
    """Render the live-AI access control without blocking the rest of the public app."""
    if not ai_available():
        st.info("Live AI is not configured yet. The deterministic workflow remains fully usable.")
        return False

    if bool_secret("ALLOW_PUBLIC_AI", False):
        st.warning("Live AI is currently public. For a portfolio deployment, use an access code to control API spend.")
        return True

    required_code = secret("AI_DEMO_CODE")
    if not required_code:
        st.info("Live AI is access-controlled and currently disabled. Configure AI_DEMO_CODE in Streamlit Secrets to enable it.")
        return False

    if st.session_state.get("ai_unlocked", False):
        st.success("Live AI demo unlocked for this browser session.")
        if st.button("Lock live AI", key="lock_live_ai"):
            st.session_state["ai_unlocked"] = False
            st.rerun()
        return True

    st.info("The workflow is public, but paid AI calls require a demo access code to prevent unauthorized API use.")
    entered = st.text_input("Live AI demo access code", type="password", key="ai_demo_code_input")
    if st.button("Unlock live AI", key="unlock_live_ai"):
        if hmac.compare_digest(str(entered), str(required_code)):
            st.session_state["ai_unlocked"] = True
            st.rerun()
        else:
            st.error("Incorrect demo access code.")
    return False


def sidebar() -> None:
    with st.sidebar:
        st.markdown(f"### {APP_TITLE}")
        st.caption(f"Portfolio proof of concept · v{APP_VERSION}")
        st.divider()
        st.write("**AI status**")
        if ai_available():
            st.success(f"Configured · {secret('OPENAI_MODEL', 'gpt-5.6-luna')}")
        else:
            st.warning("Not configured · rule engine still works")
        st.write(f"AI calls this session: {st.session_state.get('ai_calls', 0)} / {int_secret('MAX_AI_CALLS_PER_SESSION', 3)}")
        st.divider()
        if st.button("Load fictional demo data"):
            load_demo_data()
            st.rerun()
        if st.button("Clear session data"):
            st.session_state["requests"] = []
            st.session_state["sources"] = []
            st.session_state["ai_calls"] = 0
            st.rerun()
        st.caption("Data is held in the current Streamlit session only by default. Do not use real confidential customer data in the public portfolio demo.")


def dashboard_tab() -> None:
    requests = st.session_state["requests"]
    st.subheader("Operating dashboard")
    if not requests:
        st.info("No requests yet. Add one in Request Intake or load fictional demo data.")
        return

    today = date.today()
    open_requests = [r for r in requests if r.get("status") not in {"Released", "Closed"}]
    open_items = len(open_requests)
    open_snapshots = [governance_snapshot(r) for r in open_requests]
    open_enhanced = sum(1 for g in open_snapshots if g.get("effective_path") == "Enhanced")
    active_escalations = sum(1 for g in open_snapshots if g.get("governance_status") == "Escalation required")
    due_3 = 0
    overdue = 0
    for r in open_requests:
        try:
            d = date.fromisoformat(str(r.get("deadline")))
            delta = (d - today).days
            if delta < 0:
                overdue += 1
            elif delta <= 3:
                due_3 += 1
        except Exception:
            pass

    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Open requests", open_items)
    c2.metric("Open Enhanced", open_enhanced)
    c3.metric("Active AI escalations", active_escalations)
    c4.metric("Open due ≤ 3 days", due_3)
    c5.metric("Open overdue", overdue)

    display_rows = []
    for r in requests:
        g = governance_snapshot(r)
        display_rows.append({
            "request_id": r.get("request_id"),
            "customer": r.get("customer"),
            "request_type": r.get("request_type"),
            "files": len(r.get("uploaded_files", [])),
            "requirements": len(r.get("requirements", [])),
            "deadline": r.get("deadline"),
            "baseline_path": g.get("baseline_path"),
            "ai_suggested_path": g.get("ai_path") or "—",
            "ai_risk": g.get("risk_level") or "—",
            "governance_status": g.get("governance_status"),
            "urgency": r.get("urgency"),
            "status": r.get("status"),
            "owner": r.get("owner"),
        })
    display_df = pd.DataFrame(display_rows)
    st.dataframe(display_df, use_container_width=True, hide_index=True)
    st.caption("The deterministic baseline is preserved as an auditable control. AI recommendations are displayed separately and never overwrite it automatically.")

    export_df = pd.DataFrame(requests)
    st.download_button(
        "Export queue as CSV",
        export_df.to_csv(index=False).encode("utf-8"),
        file_name="deal_desk_queue.csv",
        mime="text/csv",
    )


def intake_tab() -> None:
    st.subheader("1 · Structured request intake + document package")
    st.caption(
        "Capture the minimum routing facts and optionally attach a fictional RFP / questionnaire / requirements package. "
        "Supported: PDF, DOCX, XLSX, CSV and TXT. Files are parsed in memory and raw uploads are not persisted by this demo."
    )
    with st.form("intake_form", clear_on_submit=True):
        c1, c2 = st.columns(2)
        customer = c1.text_input("Customer / requesting organization", placeholder="Fictional customer")
        entity = c2.text_input("Company legal entity / business unit", placeholder="US / UK / EU / Federal entity")
        jurisdiction = c1.text_input("Jurisdiction", placeholder="European Union")
        request_type = c2.selectbox("Request type", list(SME_MAP.keys()))
        deadline = c1.date_input("External deadline", value=date.today() + timedelta(days=7))
        owner = c2.text_input("Primary Deal Desk owner", value="Deal Desk")
        description = st.text_area(
            "Requirement / question / scope",
            height=150,
            placeholder="Summarize the request. If a requirements file is attached, a short scope note is enough.",
        )
        uploaded_files = st.file_uploader(
            "Upload request files / requirements package (fictional demo data only)",
            type=ALLOWED_UPLOAD_TYPES,
            accept_multiple_files=True,
            help=f"Up to {MAX_UPLOAD_FILES} files, {MAX_UPLOAD_FILE_BYTES // (1024*1024)} MB each. No XLS/DOC/macros. Scanned PDFs may not contain extractable text.",
        )
        consent = st.checkbox(
            "I confirm this demo input and all uploaded files contain no real confidential, privileged, personal or customer-sensitive information."
        )
        submitted = st.form_submit_button("Parse, triage and add to queue", type="primary")

    if submitted:
        uploaded_files = uploaded_files or []
        if len(uploaded_files) > MAX_UPLOAD_FILES:
            st.error(f"Upload no more than {MAX_UPLOAD_FILES} files in the public demo.")
            return
        if not description.strip() and not uploaded_files:
            st.error("Add a requirement/scope note or attach at least one supported request file.")
            return
        if not consent:
            st.error("Confirm the data-minimization statement before adding the request.")
            return

        parsed_files = []
        all_requirements: List[Dict[str, Any]] = []
        attachment_parts = []
        parse_warnings = []
        for f in uploaded_files:
            parsed = extract_uploaded_document(f)
            parsed_files.append({
                "name": parsed.get("name"),
                "extension": parsed.get("extension"),
                "size_bytes": parsed.get("size_bytes"),
                "requirement_count": len(parsed.get("requirements", [])),
                "warnings": parsed.get("warnings", []),
            })
            if parsed.get("text"):
                attachment_parts.append(f"[File: {parsed.get('name')}]\n{parsed.get('text')}")
            all_requirements.extend(parsed.get("requirements", []))
            parse_warnings.extend([f"{parsed.get('name')}: {w}" for w in parsed.get("warnings", [])])

        # De-duplicate requirement text while keeping source traceability.
        deduped = []
        seen = set()
        for req in all_requirements:
            key = normalize_space(req.get("requirement", "")).lower()
            if not key or key in seen:
                continue
            seen.add(key)
            deduped.append(req)
            if len(deduped) >= MAX_REQUIREMENTS_PER_REQUEST:
                break

        attachment_text = truncate("\n\n".join(attachment_parts), MAX_ATTACHMENT_TEXT_CHARS)
        triage_text = "\n".join([description, attachment_text[:12000]])
        triage = deterministic_triage(request_type, triage_text, deadline)
        item = {
            "request_id": f"REQ-{uuid.uuid4().hex[:8].upper()}",
            "created_at": datetime.now().isoformat(timespec="minutes"),
            "customer": customer or "Fictional / not specified",
            "entity": entity or "Not specified",
            "jurisdiction": jurisdiction or "Not specified",
            "request_type": request_type,
            "deadline": str(deadline),
            "owner": owner or "Deal Desk",
            "status": "Intake",
            "description": description,
            "uploaded_files": parsed_files,
            "attachment_text": attachment_text,
            "requirements": deduped,
            **triage,
            "ai_result": None,
            "qa": {},
            "human_override": {},
        }
        st.session_state["requests"].append(item)
        st.success(
            f"Added {item['request_id']} · {item['review_path']} review · {item['urgency']} urgency · "
            f"{len(parsed_files)} file(s) · {len(deduped)} extracted requirement(s)"
        )
        st.write("Suggested owners:", ", ".join(item["suggested_owners"]))
        if item["risk_cues"]:
            st.write("Risk cues:", ", ".join(item["risk_cues"]))
        for warning in parse_warnings:
            st.warning(warning)
        if len(all_requirements) > MAX_REQUIREMENTS_PER_REQUEST:
            st.warning(f"Requirement extraction was capped at {MAX_REQUIREMENTS_PER_REQUEST} items for the public demo.")


def requirements_tab() -> None:
    st.subheader("2 · Requirements register")
    st.caption(
        "Uploaded customer/request documents are converted into a working requirements list. "
        "Category and owner suggestions are deterministic; source matches are candidates only and require human validation."
    )
    requests = st.session_state["requests"]
    if not requests:
        st.info("Add a request first. Attach an XLSX/CSV/PDF/DOCX/TXT file in Request Intake to populate this register automatically.")
        return

    req_id = st.selectbox("Select request", [r["request_id"] for r in requests], key="requirements_req")
    req = next(r for r in requests if r["request_id"] == req_id)
    files = req.get("uploaded_files", [])
    requirements = req.get("requirements", [])

    if files:
        file_rows = []
        for f in files:
            file_rows.append({
                "file": f.get("name"),
                "type": str(f.get("extension", "")).upper(),
                "size_kb": round((f.get("size_bytes", 0) or 0) / 1024, 1),
                "extracted_requirements": f.get("requirement_count", 0),
                "parser_status": "Warning" if f.get("warnings") else "Parsed",
            })
        st.dataframe(pd.DataFrame(file_rows), use_container_width=True, hide_index=True)
        for f in files:
            for w in f.get("warnings", []) or []:
                st.warning(f"{f.get('name')}: {w}")
    else:
        st.info("This request has no uploaded files. The manually entered scope remains available for triage and AI review.")

    if not requirements:
        st.info("No structured requirements were extracted from the uploaded package.")
        return

    rows = []
    approved_sources = st.session_state["sources"]
    for r in requirements:
        source_id, score = best_candidate_source(r.get("requirement", ""), approved_sources)
        rows.append({
            "requirement_id": r.get("requirement_id"),
            "requirement": r.get("requirement"),
            "category": r.get("category"),
            "suggested_owner": r.get("suggested_owner"),
            "candidate_source": source_id,
            "match_status": "Candidate source — validate" if source_id != "—" else "No candidate source",
            "source_file": r.get("source_file"),
            "location": r.get("source_location"),
        })

    rdf = pd.DataFrame(rows)
    st.dataframe(rdf, use_container_width=True, hide_index=True)
    st.caption(
        "A candidate source is only a retrieval hint. It does not prove that the source fully answers the requirement or that its scope/entity/date is correct."
    )
    st.download_button(
        "Export requirements register as CSV",
        rdf.to_csv(index=False).encode("utf-8"),
        file_name=f"{req_id.lower()}_requirements.csv",
        mime="text/csv",
    )

def source_library_tab() -> None:
    st.subheader("3 · Canonical source library")
    st.caption("Only Approved sources are eligible for AI drafting. Draft or Deprecated material remains visible but is excluded from retrieval.")
    with st.form("source_form", clear_on_submit=True):
        c1, c2, c3 = st.columns(3)
        title = c1.text_input("Source title")
        owner = c2.text_input("Accountable source owner", placeholder="Security / GRC")
        status = c3.selectbox("Status", ["Approved", "Draft", "Deprecated"])
        review_date = c1.date_input("Next review date", value=date.today() + timedelta(days=90))
        source_id = c2.text_input("Source ID", value=f"SRC-{uuid.uuid4().hex[:6].upper()}")
        content = st.text_area("Approved content / policy excerpt / validated answer", height=150)
        submit = st.form_submit_button("Add source")
    if submit:
        if not title or not content:
            st.error("Source title and content are required.")
        else:
            st.session_state["sources"].append({
                "source_id": source_id,
                "title": title,
                "owner": owner or "Unassigned",
                "status": status,
                "review_date": str(review_date),
                "content": content,
            })
            st.success("Source added.")

    if st.session_state["sources"]:
        sdf = pd.DataFrame(st.session_state["sources"])
        st.dataframe(sdf[["source_id", "title", "owner", "status", "review_date"]], use_container_width=True, hide_index=True)
        st.download_button("Export source register as CSV", sdf.to_csv(index=False).encode("utf-8"), "canonical_sources.csv", "text/csv")
    else:
        st.info("No sources. Load demo data or add an approved source.")


def ai_review_tab() -> None:
    st.subheader("4 · AI-assisted routing, evidence retrieval and drafting")
    requests = st.session_state["requests"]
    if not requests:
        st.info("Add a request first.")
        return
    labels = {r["request_id"]: f"{r['request_id']} · {r['request_type']} · due {r['deadline']}" for r in requests}
    req_id = st.selectbox("Select request", list(labels.keys()), format_func=lambda x: labels[x])
    req = next(r for r in requests if r["request_id"] == req_id)

    st.write("**Rule-based triage:**", req.get("review_path"), "|", req.get("urgency"))
    st.write("**Suggested SMEs:**", ", ".join(req.get("suggested_owners", [])))
    retrieval_query = combined_request_text(req, 12000)
    relevant = retrieve_sources(retrieval_query, st.session_state["sources"])
    st.write("**Uploaded request files:**", ", ".join(f.get("name", "") for f in req.get("uploaded_files", [])) if req.get("uploaded_files") else "None")
    st.write("**Extracted requirements:**", len(req.get("requirements", [])))
    st.write("**Retrieved approved sources:**", ", ".join(s.get("source_id", "") for s in relevant) if relevant else "None")

    with st.expander("Privacy pre-check", expanded=False):
        redacted, n = redact_common_pii(combined_request_text(req, 14000))
        st.write(f"Common PII patterns detected/redacted before API transmission: **{n}**")
        st.code(redacted[:2500] or "No text", language=None)
        st.caption("Pattern-based redaction reduces exposure but is not a complete DLP system. Production use should add enterprise DLP/classification controls.")

    live_ai_unlocked = ai_access_panel()
    confirm = st.checkbox("I understand AI output is a draft and requires human validation before any external use.", key=f"confirm_{req_id}")
    if st.button("Run AI review", type="primary", disabled=not (confirm and live_ai_unlocked)):
        if not relevant:
            st.warning("No Approved source is available. Add an approved source before using AI drafting.")
            return
        try:
            with st.spinner("Running privacy-aware AI review..."):
                result = call_ai_analysis(req, relevant)
            req["ai_result"] = result
            req["human_override"] = {}
            req["status"] = "In Review"
            st.success("AI review completed. Human approval is still required.")
        except Exception as exc:
            st.error(str(exc))

    result = req.get("ai_result")
    if result:
        governance = governance_snapshot(req)
        c1, c2, c3 = st.columns(3)
        c1.markdown(f"**Deterministic baseline:** {governance.get('baseline_path') or 'Not provided'}")
        c2.markdown(f"**AI suggested path:** {governance.get('ai_path') or 'Not provided'}")
        c3.markdown(f"**AI risk level:** {governance.get('risk_level') or 'Not provided'}")

        if governance.get("path_discrepancy"):
            st.warning(
                "Escalation required: the AI recommendation differs from deterministic triage. "
                "The AI suggestion is advisory and does not overwrite the baseline automatically."
            )
        if governance.get("blockers"):
            st.error("Release blockers remain. They must be resolved or explicitly handled by an accountable human before release.")

        st.write("**Summary**")
        st.write(result.get("summary"))
        st.write("**Flags**")
        st.write(result.get("detected_flags") or [])
        st.write("**Missing information**")
        st.write(result.get("missing_information") or [])
        st.write("**Source IDs used**")
        st.write(result.get("source_ids_used") or [])
        st.write("**Draft response — NOT APPROVED**")
        st.text_area("Draft", value=result.get("draft_response", ""), height=220, key=f"draft_{req_id}")
        st.info("The app intentionally does not provide an autonomous Send/Submit action.")


def qa_tab() -> None:
    st.subheader("5 · Second-read QA and human release gate")
    requests = st.session_state["requests"]
    if not requests:
        st.info("Add a request first.")
        return
    req_id = st.selectbox("Select request for QA", [r["request_id"] for r in requests], key="qa_req")
    req = next(r for r in requests if r["request_id"] == req_id)
    governance = governance_snapshot(req)

    c1, c2, c3 = st.columns(3)
    c1.metric("Deterministic baseline", governance.get("baseline_path") or "—")
    c2.metric("AI suggested path", governance.get("ai_path") or "Not run")
    c3.metric("AI risk", governance.get("risk_level") or "Not run")

    if governance.get("path_discrepancy"):
        st.warning(
            "AI escalation detected. The deterministic baseline remains unchanged for auditability; "
            "the higher AI recommendation must be reviewed by a human."
        )

    blockers = governance.get("blockers") or []
    if blockers:
        st.error("HOLD: unresolved AI-governance blockers prevent normal release.")
        for blocker in blockers:
            st.write(f"• {blocker}")

    checks = [
        "Scope, customer, legal entity, jurisdiction and package version confirmed",
        "Deadline and portal/submission mechanics confirmed",
        "Uploaded request package and extracted requirements reviewed for completeness",
        "All required fields and attachments accounted for",
        "Material claims trace to Approved canonical sources",
        "Entity and certification distinctions verified",
        "Legal / Privacy / Security / Finance / Product approvals obtained where required",
        "Contradictory or stale sources resolved by accountable owner",
        "A second reader reviewed names, dates, figures, attachments and consistency",
        "Named final approver confirmed readiness",
        "Submission evidence / receipt will be retained after release",
    ]
    qa = req.setdefault("qa", {})
    for idx, check in enumerate(checks):
        qa[str(idx)] = st.checkbox(check, value=qa.get(str(idx), False), key=f"qa_{req_id}_{idx}")

    approved = all(qa.values()) if qa else False
    override = req.setdefault("human_override", {})

    if blockers:
        with st.expander("Human resolution / exception record", expanded=not bool(override.get("approved"))):
            st.caption(
                "This does not let AI approve its own output. An accountable human must document how the blockers "
                "were resolved or why an exception is being accepted for this fictional demo."
            )
            approver = st.text_input(
                "Named accountable approver",
                value=override.get("approver", ""),
                key=f"override_approver_{req_id}",
            )
            decision = st.selectbox(
                "Human decision",
                ["Issues resolved after human review", "Exception accepted for fictional demo release"],
                index=0 if override.get("decision") != "Exception accepted for fictional demo release" else 1,
                key=f"override_decision_{req_id}",
            )
            rationale = st.text_area(
                "Resolution / exception rationale",
                value=override.get("rationale", ""),
                height=120,
                placeholder="Describe what was resolved, what evidence was reviewed, or why the exception is acceptable (minimum 30 characters).",
                key=f"override_rationale_{req_id}",
            )
            rationale_len = len(rationale.strip())
            st.caption(f"Rationale length: {rationale_len}/30 minimum characters")

            exception_ack = True
            if decision == "Exception accepted for fictional demo release":
                st.warning(
                    "An exception does not resolve the underlying AI flags or missing evidence. "
                    "It records an accountable human decision to proceed in this fictional demo despite unresolved items."
                )
                exception_ack = st.checkbox(
                    "I understand this exception does not mean the missing evidence or AI flags were resolved.",
                    key=f"override_exception_ack_{req_id}",
                )

            attested = st.checkbox(
                "I confirm an accountable human reviewed the AI flags, missing information and triage discrepancy before this decision.",
                key=f"override_attest_{req_id}",
            )
            if st.button("Record human decision", key=f"record_override_{req_id}"):
                if not approver.strip():
                    st.error("A named accountable approver is required.")
                elif rationale_len < 30:
                    st.error("Provide a substantive rationale of at least 30 characters.")
                elif not attested:
                    st.error("Human-review confirmation is required.")
                elif not exception_ack:
                    st.error("Confirm that the exception does not resolve the underlying evidence gaps.")
                else:
                    req["human_override"] = {
                        "approved": True,
                        "approver": approver.strip(),
                        "decision": decision,
                        "rationale": rationale.strip(),
                        "recorded_at": datetime.now().isoformat(timespec="minutes"),
                        "blockers_at_decision": list(blockers),
                    }
                    st.success("Human decision recorded. The release gate will still require every QA check.")
                    st.rerun()

    override_ok = bool(req.get("human_override", {}).get("approved"))
    release_ready = approved and (not blockers or override_ok)

    if release_ready:
        if blockers:
            recorded = req.get("human_override", {})
            st.success(
                f"QA gate passed with documented human decision by {recorded.get('approver')}. "
                "A human final approver may release the submission outside this demo."
            )
        else:
            st.success("QA gate passed. A human final approver may release the submission outside this demo.")
        if st.button("Mark as Released (demo status only)", key=f"release_{req_id}"):
            req["status"] = "Released"
            req["released_at"] = datetime.now().isoformat(timespec="minutes")
            st.rerun()
    elif approved and blockers and not override_ok:
        st.error("HOLD: QA checks are complete, but AI-governance blockers still require a documented human decision.")
    else:
        st.warning("HOLD: release gate is not complete.")


def architecture_tab() -> None:
    st.subheader("6 · Why the controls exist")
    st.markdown(
        """
**Design principle:** AI should create leverage around repetitive, information-intensive work while accountable humans retain final authority over externally material claims and commitments.

- **Structured intake + document ingestion** accepts fictional PDF/DOCX/XLSX/CSV/TXT request packages, extracts a requirements register locally, and preserves file/row/page traceability without treating customer documents as company evidence.
- **Deterministic triage** creates a transparent baseline that does not depend on model output; AI may recommend escalation, but it never silently overwrites that baseline.
- **Canonical-source library** prevents plausible-but-unverified answers from becoming customer-facing commitments.
- **AI retrieval and drafting** accelerates research and first drafts, but only from Approved material.
- **PII redaction + data minimization** reduces unnecessary personal-data exposure before API calls.
- **SME routing** makes the true accountable owner visible early enough to unblock the request.
- **Second-read QA** detects contradictions, missing evidence and wrong entity/certification details before release.
- **Human release gate** blocks release when AI surfaces unresolved High-risk issues, missing information or a triage discrepancy; any exception requires a named human decision, a substantive rationale, and explicit acknowledgement that unresolved evidence gaps remain.
- **Session-only storage by default** makes the public portfolio demo safer; production systems should use authenticated, access-controlled, audited storage.
        """
    )
    st.warning("Portfolio demo only. Do not use this deployment for confidential customer, legal, security, HR or personal data.")


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="⚙️", layout="wide")
    init_state()
    sidebar()

    st.title(APP_TITLE)
    st.write("A fictional, human-in-the-loop proof of concept for structured intake, document/requirements ingestion, risk-based routing, canonical-source governance, AI-assisted drafting, deadline control and release QA.")
    st.caption("All demo content is fictional. The application deliberately separates AI assistance from accountable human approval.")

    tabs = st.tabs(["Dashboard", "Request Intake", "Requirements", "Source Library", "AI Review", "QA & Release", "Architecture"])

    # Render form-driven tabs first so newly submitted data is available
    # to the Dashboard during the same Streamlit rerun. The visual tab
    # order remains unchanged.
    with tabs[1]: intake_tab()
    with tabs[3]: source_library_tab()
    with tabs[2]: requirements_tab()
    with tabs[4]: ai_review_tab()
    with tabs[5]: qa_tab()
    with tabs[6]: architecture_tab()
    with tabs[0]: dashboard_tab()


if __name__ == "__main__":
    main()
